"""Notes endpoints: OCR draft -> review -> save, typed and uploaded notes, downloads, practice quiz."""

import io
import json

import pymupdf
import pytest
from docx import Document
from PIL import Image
from sqlalchemy import delete

from app.ai import AIError
from app.models import Note, NotePage, NoteShare
from tests.conftest import login
from tests.test_notes_access import AARAV, TEACHER, png


@pytest.fixture(autouse=True)
def clean_notes(db):
    yield
    db.execute(delete(NoteShare))
    db.execute(delete(NotePage))
    db.execute(delete(Note))
    db.commit()


def ocr(client, who, files, replies, provider, **data):
    provider.replies += [json.dumps({"text": r}) for r in replies]
    return client.post("/api/notes/ocr", files=files, data=data, headers=login(client, who))


# ---------------------------------------------------------------- OCR flow


def test_ocr_creates_a_draft_with_page_images(client, provider, storage):
    resp = ocr(client, AARAV, [("files", ("a.png", png(), "image/png"))], ["Magnets\n- A [?magnet?] attracts iron"], provider)
    assert resp.status_code == 201
    note = resp.json()
    assert (note["status"], note["source_type"], note["page_count"]) == ("draft", "ocr", 1)
    assert note["title"] == "Magnets"
    assert "[?magnet?]" in note["text"]  # markers kept so the review screen can highlight them
    # the original page is viewable by its owner
    img = client.get(f"/api/notes/{note['id']}/pages/1/image", headers=login(client, AARAV))
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"
    # OCR asks the vision model, once per page, deterministically
    call = provider.calls[0]
    assert call["temperature"] == 0.0 and isinstance(call["messages"][1]["content"], list)


def test_drafts_stay_out_of_the_library_until_saved(client, provider):
    note = ocr(client, AARAV, [("files", ("a.png", png(), "image/png"))], ["Heading\n- one [?two?]"], provider).json()
    headers = login(client, AARAV)
    assert client.get("/api/notes", headers=headers).json() == []

    saved = client.patch(
        f"/api/notes/{note['id']}", json={"title": "My magnets", "text": "Heading\n- one two", "status": "saved"}, headers=headers
    )
    assert saved.status_code == 200 and saved.json()["status"] == "saved"
    assert [n["title"] for n in client.get("/api/notes", headers=headers).json()] == ["My magnets"]


def test_saving_strips_any_leftover_uncertain_markers(client, provider):
    note = ocr(client, AARAV, [("files", ("a.png", png(), "image/png"))], ["- [?iron?] and [?illegible?]"], provider).json()
    saved = client.patch(f"/api/notes/{note['id']}", json={"text": note["text"], "status": "saved"}, headers=login(client, AARAV)).json()
    assert saved["text"] == "- iron and [illegible]"


def test_multi_page_notes_are_read_in_order_and_joined(client, provider):
    files = [("files", (f"p{i}.png", png(), "image/png")) for i in (1, 2)]
    note = ocr(client, TEACHER, files, ["First page", "Second page"], provider).json()
    assert note["page_count"] == 2
    assert note["text"] == "First page\n\nSecond page"
    assert len(provider.calls) == 2


def test_pdf_pages_are_rendered_and_read(client, provider):
    buf = io.BytesIO()
    a, b = Image.new("RGB", (100, 140), "white"), Image.new("RGB", (100, 140), "white")
    a.save(buf, "PDF", save_all=True, append_images=[b])
    note = ocr(client, TEACHER, [("files", ("scan.pdf", buf.getvalue(), "application/pdf"))], ["one", "two"], provider).json()
    assert note["page_count"] == 2


def test_too_many_pages_is_refused_before_any_ai_call(client, provider):
    files = [("files", (f"p{i}.png", png(), "image/png")) for i in range(6)]
    resp = client.post("/api/notes/ocr", files=files, headers=login(client, AARAV))
    assert resp.status_code == 422 and "5 pages" in resp.json()["detail"]
    assert provider.calls == []


def test_bad_uploads_get_friendly_errors(client, provider):
    headers = login(client, AARAV)
    not_image = client.post("/api/notes/ocr", files=[("files", ("x.png", b"not an image", "image/png"))], headers=headers)
    assert not_image.status_code == 422 and "valid image" in not_image.json()["detail"]
    wrong_type = client.post("/api/notes/ocr", files=[("files", ("x.exe", b"MZ", "application/octet-stream"))], headers=headers)
    assert wrong_type.status_code == 422 and "JPG, PNG or PDF" in wrong_type.json()["detail"]
    empty = client.post("/api/notes/ocr", files=[("files", ("x.png", b"", "image/png"))], headers=headers)
    assert empty.status_code == 422
    assert provider.calls == []


def test_blank_page_is_reported_not_saved(client, provider):
    resp = ocr(client, AARAV, [("files", ("a.png", png(), "image/png"))], [""], provider)
    assert resp.status_code == 422
    assert client.get("/api/notes", headers=login(client, AARAV)).json() == []


def test_ai_failure_becomes_a_friendly_503_and_leaves_nothing_behind(client, provider):
    provider.replies.append(AIError("The AI service is busy (rate limit).", retryable=False))
    resp = client.post("/api/notes/ocr", files=[("files", ("a.png", png(), "image/png"))], headers=login(client, AARAV))
    assert resp.status_code == 503 and "busy" in resp.json()["detail"]
    assert client.get("/api/notes", headers=login(client, AARAV)).json() == []


# ---------------------------------------------------------------- typed and uploaded notes


def test_typed_note_is_saved_and_searchable(client):
    headers = login(client, AARAV)
    made = client.post("/api/notes", json={"title": "Water", "text": "Ice melts at 0 degrees."}, headers=headers)
    assert made.status_code == 201 and made.json()["source_type"] == "typed"
    hits = client.get("/api/notes", params={"q": "MELTS"}, headers=headers).json()  # case-insensitive
    assert [n["title"] for n in hits] == ["Water"]
    assert client.get("/api/notes", params={"q": "volcano"}, headers=headers).json() == []


def test_search_treats_percent_and_underscore_literally(client):
    headers = login(client, AARAV)
    client.post("/api/notes", json={"title": "Plain", "text": "nothing special"}, headers=headers)
    assert client.get("/api/notes", params={"q": "%"}, headers=headers).json() == []
    assert client.get("/api/notes", params={"q": "_"}, headers=headers).json() == []


def test_txt_docx_and_typed_pdf_uploads(client):
    headers = login(client, TEACHER)
    txt = client.post("/api/notes/upload", files={"file": ("moon.txt", b"The Moon goes around the Earth.", "text/plain")}, headers=headers)
    assert txt.status_code == 201 and txt.json()["title"] == "moon"

    buf = io.BytesIO()
    d = Document()
    d.add_paragraph("Iron is magnetic.")
    d.save(buf)
    docx = client.post("/api/notes/upload", files={"file": ("m.docx", buf.getvalue(), "application/octet-stream")}, data={"title": "Iron"}, headers=headers)
    assert docx.json()["title"] == "Iron" and "Iron is magnetic." in docx.json()["text"]

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Winnowing separates husk from grain using wind. " * 4)
    typed = client.post("/api/notes/upload", files={"file": ("w.pdf", pdf.tobytes(), "application/pdf")}, headers=headers)
    assert typed.status_code == 201 and "Winnowing" in typed.json()["text"]


def test_scanned_pdf_upload_points_to_handwriting_digitiser(client):
    buf = io.BytesIO()
    Image.new("RGB", (200, 200), "white").save(buf, "PDF")
    resp = client.post("/api/notes/upload", files={"file": ("scan.pdf", buf.getvalue(), "application/pdf")}, headers=login(client, AARAV))
    assert resp.status_code == 422 and "Digitise handwriting" in resp.json()["detail"]


def test_wrong_file_type_upload_is_refused(client):
    resp = client.post("/api/notes/upload", files={"file": ("x.exe", b"MZ", "application/octet-stream")}, headers=login(client, AARAV))
    assert resp.status_code == 422


def test_unknown_chapter_is_refused(client):
    resp = client.post("/api/notes", json={"title": "x", "text": "y", "chapter_id": 9999}, headers=login(client, AARAV))
    assert resp.status_code == 422


# ---------------------------------------------------------------- edit, download


def test_owner_can_edit_and_cannot_empty_a_note(client):
    headers = login(client, AARAV)
    note_id = client.post("/api/notes", json={"title": "T", "text": "body"}, headers=headers).json()["id"]
    edited = client.patch(f"/api/notes/{note_id}", json={"title": "New", "topic": "Poles"}, headers=headers).json()
    assert (edited["title"], edited["topic"], edited["text"]) == ("New", "Poles", "body")
    assert client.patch(f"/api/notes/{note_id}", json={"text": "  "}, headers=headers).status_code == 422
    assert client.patch(f"/api/notes/{note_id}", json={"status": "draft"}, headers=headers).status_code == 422


def test_docx_download_has_title_bullets_and_a_filename(client):
    headers = login(client, AARAV)
    note_id = client.post("/api/notes", json={"title": "Magnets & Poles!", "text": "Intro line\n- first\n- second"}, headers=headers).json()["id"]
    resp = client.get(f"/api/notes/{note_id}/download", headers=headers)
    assert resp.status_code == 200
    assert 'filename="Magnets_Poles.docx"' in resp.headers["content-disposition"]
    doc = Document(io.BytesIO(resp.content))
    assert [p.text for p in doc.paragraphs] == ["Magnets & Poles!", "Intro line", "first", "second"]
    assert doc.paragraphs[2].style.name == "List Bullet"


# ---------------------------------------------------------------- practice quiz

QUIZ = {
    "questions": [
        {"question": "What do magnets attract?", "options": ["Iron", "Wood", "Glass", "Paper"], "correct_index": 0, "explanation": "Magnets attract iron."},
        {"question": "Magnets have two poles.", "options": ["True", "False"], "correct_index": 0, "explanation": "North and south."},
    ]
}
LONG_TEXT = "Magnets attract iron, nickel and cobalt. " * 10


def test_student_practice_quiz(client, provider):
    headers = login(client, AARAV)
    note_id = client.post("/api/notes", json={"title": "M", "text": LONG_TEXT}, headers=headers).json()["id"]
    provider.replies.append(json.dumps(QUIZ))
    resp = client.post(f"/api/notes/{note_id}/practice", headers=headers)
    assert resp.status_code == 200
    assert [q["correct_index"] for q in resp.json()["questions"]] == [0, 0]
    assert provider.calls[0]["model"] == "openai/gpt-oss-20b"  # cheap fast model


def test_practice_quiz_rejects_malformed_model_output(client, provider):
    headers = login(client, AARAV)
    note_id = client.post("/api/notes", json={"title": "M", "text": LONG_TEXT}, headers=headers).json()["id"]
    bad = {"questions": [{"question": "What is it?", "options": ["a", "b", "c", "d"], "correct_index": 7, "explanation": "x"}]}
    provider.replies += [json.dumps(bad)] * 3
    resp = client.post(f"/api/notes/{note_id}/practice", headers=headers)
    assert resp.status_code == 503  # invalid output is retried, then reported, never shown to a student


def test_practice_quiz_is_students_only_and_needs_enough_text(client):
    teacher = login(client, TEACHER)
    note_id = client.post("/api/notes", json={"title": "M", "text": LONG_TEXT}, headers=teacher).json()["id"]
    assert client.post(f"/api/notes/{note_id}/practice", headers=teacher).status_code == 403
    student = login(client, AARAV)
    short_id = client.post("/api/notes", json={"title": "S", "text": "Too short."}, headers=student).json()["id"]
    resp = client.post(f"/api/notes/{short_id}/practice", headers=student)
    assert resp.status_code == 422 and "too short" in resp.json()["detail"]


def test_chapters_and_classes_lists(client):
    chapters = client.get("/api/chapters", headers=login(client, AARAV)).json()
    assert len(chapters) == 12 and chapters[0]["number"] == 1
    assert client.get("/api/classes", headers=login(client, TEACHER)).json()[0]["name"] == "6-A"
    assert client.get("/api/classes", headers=login(client, AARAV)).status_code == 403
    assert client.get("/api/chapters").status_code == 401
