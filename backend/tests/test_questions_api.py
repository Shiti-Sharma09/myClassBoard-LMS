"""Question Bank API: sets and versions, review, the bank, papers, assessments, and who may do what."""

import io

import pytest
from docx import Document
from sqlalchemy import delete, select

from app.ai import AIError
from app.models import (
    Assessment,
    AssessmentQuestion,
    Chapter,
    Note,
    NotePage,
    NoteShare,
    Question,
    QuestionSet,
    SchoolClass,
    User,
)
from app.security import hash_password
from tests.conftest import login
from tests.qhelpers import NOTE_TEXT, POOL, pool, reply

TEACHER, AARAV = "teacher@demo.school", "aarav@demo.school"
ALL_TYPES = ["mcq", "true_false", "fill_blank", "short", "long"]


@pytest.fixture(autouse=True)
def clean(db):
    yield
    for table in (AssessmentQuestion, Assessment, Question, QuestionSet, NoteShare, NotePage, Note):
        db.execute(delete(table))
    db.execute(delete(User).where(User.email == "other.teacher@demo.school"))
    db.commit()


@pytest.fixture
def note_id(client):
    resp = client.post("/api/notes", json={"title": "Exploring Magnets", "text": NOTE_TEXT}, headers=login(client, TEACHER))
    assert resp.status_code == 201
    return resp.json()["id"]


def teacher(client):
    return login(client, TEACHER)


def make_set(client, provider, note_id, items=None, count=5, **body):
    provider.replies.append(reply(items if items is not None else pool(0, 8)))
    payload = {"count": count, "types": ALL_TYPES, "difficulty": "mixed", **body}
    return client.post(f"/api/notes/{note_id}/question-sets", json=payload, headers=teacher(client))


def accept_all(client, set_id):
    return client.post(f"/api/question-sets/{set_id}/bulk", json={"status": "accepted"}, headers=teacher(client))


# ---------------------------------------------------------------- generating


def test_first_version_is_balanced_with_draft_questions(client, provider, note_id):
    resp = make_set(client, provider, note_id)
    assert resp.status_code == 201, resp.text
    qs = resp.json()
    assert (qs["version"], qs["emphasis"], qs["requested_count"]) == (1, "balanced", 5)
    assert len(qs["questions"]) == 5 and qs["shortfall_message"] is None
    assert {q["status"] for q in qs["questions"]} == {"draft"}
    assert qs["generation_seconds"] is not None
    marks = {q["type"]: q["marks"] for q in qs["questions"]}
    assert marks.get("mcq", 1) == 1 and marks.get("long", 5) == 5  # sensible default marks by type
    assert [q["position"] for q in qs["questions"]] == list(range(5))


def test_versions_increase_and_cap_at_three(client, provider, note_id):
    versions = []
    for start in (0, 4, 8):
        resp = make_set(client, provider, note_id, items=pool(start, start + 6), count=3)
        assert resp.status_code == 201, resp.text
        versions.append((resp.json()["version"], resp.json()["emphasis"]))
    assert versions == [(1, "balanced"), (2, "application"), (3, "recall")]

    provider.replies.append(reply(pool(0, 8)))
    fourth = client.post(f"/api/notes/{note_id}/question-sets", json={"count": 3, "types": ALL_TYPES, "difficulty": "easy"}, headers=teacher(client))
    assert fourth.status_code == 409 and "3 versions" in fourth.json()["detail"]
    assert len(provider.replies) == 1  # refused before wasting an AI call


def test_discarding_a_version_frees_its_slot(client, provider, note_id):
    ids = [make_set(client, provider, note_id, items=pool(s, s + 6), count=3).json()["id"] for s in (0, 4, 8)]
    assert client.delete(f"/api/question-sets/{ids[1]}", headers=teacher(client)).status_code == 204
    again = make_set(client, provider, note_id, items=pool(1, 8), count=3)
    assert again.status_code == 201 and again.json()["version"] == 2


def test_a_version_with_accepted_questions_cannot_be_deleted(client, provider, note_id):
    set_id = make_set(client, provider, note_id).json()["id"]
    accept_all(client, set_id)
    resp = client.delete(f"/api/question-sets/{set_id}", headers=teacher(client))
    assert resp.status_code == 409


def test_later_versions_avoid_earlier_questions(client, provider, note_id):
    first = make_set(client, provider, note_id, items=pool(0, 6), count=4).json()
    second = make_set(client, provider, note_id, items=pool(0, 12), count=4).json()  # model repeats the first six
    assert second["version"] == 2
    first_texts = {q["text"] for q in first["questions"]}
    assert first_texts.isdisjoint({q["text"] for q in second["questions"]})


def test_shortfall_is_reported_and_stored(client, provider, note_id):
    provider.replies += [reply(pool(0, 2)), reply(pool(0, 2))]
    resp = client.post(f"/api/notes/{note_id}/question-sets", json={"count": 6, "types": ALL_TYPES, "difficulty": "mixed"}, headers=teacher(client))
    assert resp.status_code == 201
    assert len(resp.json()["questions"]) == 2 and "only supported 2" in resp.json()["shortfall_message"]
    listed = client.get(f"/api/notes/{note_id}/question-sets", headers=teacher(client)).json()
    assert listed[0]["shortfall_message"] == resp.json()["shortfall_message"]


def test_nothing_usable_is_a_clear_error_and_nothing_is_stored(client, provider, note_id):
    provider.replies.append(reply([], insufficient=True))
    resp = client.post(f"/api/notes/{note_id}/question-sets", json={"count": 5, "types": ALL_TYPES, "difficulty": "mixed"}, headers=teacher(client))
    assert resp.status_code == 422 and "No good questions" in resp.json()["detail"]
    assert client.get(f"/api/notes/{note_id}/question-sets", headers=teacher(client)).json() == []


def test_ai_failure_is_a_friendly_503_and_stores_nothing(client, provider, note_id):
    provider.replies.append(AIError("The AI service is busy (rate limit)."))
    resp = client.post(f"/api/notes/{note_id}/question-sets", json={"count": 5, "types": ALL_TYPES, "difficulty": "mixed"}, headers=teacher(client))
    assert resp.status_code == 503 and "busy" in resp.json()["detail"]
    assert client.get(f"/api/notes/{note_id}/question-sets", headers=teacher(client)).json() == []


@pytest.mark.parametrize(
    "body",
    [
        {"count": 0, "types": ["mcq"]},
        {"count": 21, "types": ["mcq"]},
        {"count": 5, "types": []},
        {"count": 5, "types": ["riddle"]},
        {"count": 5, "types": ["mcq"], "difficulty": "impossible"},
    ],
)
def test_bad_generation_requests_are_refused_before_any_ai_call(client, provider, note_id, body):
    resp = client.post(f"/api/notes/{note_id}/question-sets", json=body, headers=teacher(client))
    assert resp.status_code == 422 and provider.calls == []


def test_topics_are_found_without_an_ai_call_and_can_be_chosen(client, provider, note_id):
    topics = client.get(f"/api/notes/{note_id}/topics", headers=teacher(client)).json()
    assert [t["title"] for t in topics] == ["Magnetic Materials", "Poles and Directions", "Making and Keeping Magnets", "Quick Recap"]
    assert [t["recap"] for t in topics] == [False, False, False, True]
    assert provider.calls == []

    resp = make_set(client, provider, note_id, section_indices=[1])
    assert resp.status_code == 201
    prompt = provider.calls[0]["messages"][1]["content"]
    assert "### Poles and Directions" in prompt and "### Magnetic Materials" not in prompt
    bad = client.post(f"/api/notes/{note_id}/question-sets", json={"count": 3, "types": ["mcq"], "difficulty": "easy", "section_indices": [99]}, headers=teacher(client))
    assert bad.status_code == 422


def test_a_draft_note_cannot_have_questions_generated(client, db):
    owner = db.scalar(select(User.id).where(User.email == TEACHER))
    note = Note(owner_id=owner, title="Draft", source_type="ocr", status="draft", text="Some handwriting text here.")
    db.add(note)
    db.commit()
    resp = client.post(f"/api/notes/{note.id}/question-sets", json={"count": 3, "types": ["mcq"], "difficulty": "easy"}, headers=teacher(client))
    assert resp.status_code == 422 and "Save the note" in resp.json()["detail"]


# ---------------------------------------------------------------- who may do what


def test_only_teachers_can_use_the_question_bank(client, provider, note_id):
    for who in (AARAV, "parent.aarav@demo.school", "admin@demo.school"):
        headers = login(client, who)
        assert client.get("/api/question-bank", headers=headers).status_code == 403
        assert client.post(f"/api/notes/{note_id}/question-sets", json={"count": 3, "types": ["mcq"], "difficulty": "easy"}, headers=headers).status_code == 403
    assert client.get("/api/question-bank").status_code == 401
    assert provider.calls == []


def test_another_teacher_cannot_see_or_change_my_questions(client, provider, note_id, db):
    qs = make_set(client, provider, note_id).json()
    accept_all(client, qs["id"])
    db.add(User(name="Other", email="other.teacher@demo.school", role="teacher", password_hash=hash_password("demo1234")))
    db.commit()
    other = login(client, "other.teacher@demo.school")
    qid = qs["questions"][0]["id"]
    assert client.get(f"/api/notes/{note_id}/question-sets", headers=other).status_code == 404
    assert client.post(f"/api/notes/{note_id}/question-sets", json={"count": 3, "types": ["mcq"], "difficulty": "easy"}, headers=other).status_code == 404
    assert client.patch(f"/api/questions/{qid}", json={"status": "discarded"}, headers=other).status_code == 404
    assert client.delete(f"/api/question-sets/{qs['id']}", headers=other).status_code == 404
    assert client.post(f"/api/question-sets/{qs['id']}/bulk", json={"status": "discarded"}, headers=other).status_code == 404
    assert client.get("/api/question-bank", headers=other).json() == []
    paper = {"title": "x", "sections": [{"title": "A", "questions": [{"id": qid, "marks": 1}]}]}
    assert client.post("/api/papers/docx", json=paper, headers=other).status_code == 404


# ---------------------------------------------------------------- review


def test_accept_discard_and_edit(client, provider, note_id):
    qs = make_set(client, provider, note_id).json()
    q = next(x for x in qs["questions"] if x["type"] == "mcq")
    headers = teacher(client)

    accepted = client.patch(f"/api/questions/{q['id']}", json={"status": "accepted"}, headers=headers).json()
    assert accepted["status"] == "accepted" and accepted["edited"] is False  # accepting alone is not an edit

    edited = client.patch(f"/api/questions/{q['id']}", json={"text": "Which material do magnets attract most strongly?", "marks": 2}, headers=headers).json()
    assert edited["edited"] is True and edited["marks"] == 2 and edited["status"] == "accepted"

    assert client.patch(f"/api/questions/{q['id']}", json={"status": "discarded"}, headers=headers).json()["status"] == "discarded"
    assert client.patch(f"/api/questions/{q['id']}", json={"status": "draft"}, headers=headers).json()["status"] == "draft"


def test_edits_that_would_break_a_question_are_refused(client, provider, note_id):
    qs = make_set(client, provider, note_id, items=pool(0, 15), count=15, types=ALL_TYPES).json()
    by_type = {}
    for x in qs["questions"]:
        by_type.setdefault(x["type"], x)
    headers = teacher(client)

    def patch(q, **body):
        return client.patch(f"/api/questions/{q['id']}", json=body, headers=headers)

    mcq, fill, tf = by_type["mcq"], by_type["fill_blank"], by_type["true_false"]
    assert "match one of the options" in patch(mcq, answer="Plutonium").json()["detail"]
    assert "4 different options" in patch(mcq, options=["a", "a", "b", "c"]).json()["detail"]
    fixed = patch(mcq, options=["Steel", "Wood", "Glass", "Rubber"], answer="steel")
    assert fixed.status_code == 200 and fixed.json()["answer"] == "Steel"
    assert "needs a blank" in patch(fill, text="Like poles repel each other always.").json()["detail"]
    assert "True or False" in patch(tf, answer="maybe").json()["detail"]
    assert patch(tf, answer="false").json()["answer"] == "False"
    assert patch(mcq, marks=-1).status_code == 422 and patch(mcq, marks=101).status_code == 422
    assert patch(mcq, difficulty="extreme").status_code == 422
    assert patch(mcq, bloom="Guess").status_code == 422
    assert patch(mcq, status="deleted").status_code == 422
    assert patch(mcq, topic="  ").status_code == 422
    assert patch(mcq, text="Hi").status_code == 422
    assert client.patch("/api/questions/999999", json={"status": "accepted"}, headers=headers).status_code == 404


def test_bulk_accept_only_touches_unreviewed_questions(client, provider, note_id):
    qs = make_set(client, provider, note_id).json()
    headers = teacher(client)
    first, second = qs["questions"][0], qs["questions"][1]
    client.patch(f"/api/questions/{first['id']}", json={"status": "discarded"}, headers=headers)
    result = client.post(f"/api/question-sets/{qs['id']}/bulk", json={"status": "accepted"}, headers=headers).json()
    statuses = {q["id"]: q["status"] for q in result["questions"]}
    assert statuses[first["id"]] == "discarded" and statuses[second["id"]] == "accepted"
    assert client.post(f"/api/question-sets/{qs['id']}/bulk", json={"status": "maybe"}, headers=headers).status_code == 422


# ---------------------------------------------------------------- the bank


def test_bank_lists_only_accepted_and_filters(client, provider, note_id):
    qs = make_set(client, provider, note_id, items=pool(0, 15), count=15).json()
    headers = teacher(client)
    assert client.get("/api/question-bank", headers=headers).json() == []  # nothing accepted yet
    accept_all(client, qs["id"])
    items = client.get("/api/question-bank", headers=headers).json()
    assert len(items) == 15 and all(i["note_title"] == "Exploring Magnets" for i in items)

    mcqs = client.get("/api/question-bank", params={"type": "mcq"}, headers=headers).json()
    assert mcqs and {i["type"] for i in mcqs} == {"mcq"}
    assert client.get("/api/question-bank", params={"difficulty": "hard"}, headers=headers).json() == []
    assert len(client.get("/api/question-bank", params={"bloom": "Remember"}, headers=headers).json()) == 15
    assert {i["topic"] for i in client.get("/api/question-bank", params={"topic": "Poles"}, headers=headers).json()} == {"Poles"}


def test_bank_tree_groups_by_subject_chapter_and_topic_with_counts(client, provider, db):
    chapter = db.scalar(select(Chapter).where(Chapter.number == 4))
    with_chapter = client.post("/api/notes", json={"title": "Magnets A", "text": NOTE_TEXT, "chapter_id": chapter.id}, headers=teacher(client)).json()["id"]
    without = client.post("/api/notes", json={"title": "Magnets B", "text": NOTE_TEXT}, headers=teacher(client)).json()["id"]
    a = make_set(client, provider, with_chapter, items=pool(0, 8), count=6).json()
    b = make_set(client, provider, without, items=pool(0, 8), count=3).json()
    accept_all(client, a["id"])
    accept_all(client, b["id"])

    tree = client.get("/api/question-bank/tree", headers=teacher(client)).json()
    assert [s["subject"] for s in tree] == ["Science"] and tree[0]["count"] == 9
    chapters = {c["chapter_title"]: c for c in tree[0]["chapters"]}
    assert chapters["Exploring Magnets"]["count"] == 6 and chapters["No chapter"]["count"] == 3
    assert sum(t["count"] for t in chapters["Exploring Magnets"]["topics"]) == 6
    assert tree[0]["chapters"][-1]["chapter_title"] == "No chapter"  # unassigned last

    only_mcq = client.get("/api/question-bank/tree", params={"type": "mcq"}, headers=teacher(client)).json()
    assert only_mcq[0]["count"] < 9

    by_chapter = client.get("/api/question-bank", params={"chapter_id": chapter.id}, headers=teacher(client)).json()
    assert len(by_chapter) == 6
    assert len(client.get("/api/question-bank", params={"no_chapter": "true"}, headers=teacher(client)).json()) == 3


def test_usable_rate_counts_accepted_over_reviewed(client, provider, note_id):
    headers = teacher(client)
    assert client.get("/api/question-bank/stats", headers=headers).json()["usable_rate"] is None
    qs = make_set(client, provider, note_id, count=5).json()
    ids = [q["id"] for q in qs["questions"]]
    for qid in ids[:4]:
        client.patch(f"/api/questions/{qid}", json={"status": "accepted"}, headers=headers)
    client.patch(f"/api/questions/{ids[4]}", json={"status": "discarded"}, headers=headers)
    client.patch(f"/api/questions/{ids[0]}", json={"text": "Which material is attracted by a magnet most?"}, headers=headers)
    stats = client.get("/api/question-bank/stats", headers=headers).json()
    assert (stats["accepted"], stats["discarded"], stats["pending"], stats["edited"]) == (4, 1, 0, 1)
    assert stats["usable_rate"] == 0.8


# ---------------------------------------------------------------- paper


def _paper_body(questions, **over):
    return {
        "school": "myClassBoard Demo School", "title": "Unit Test: Magnets", "class_name": "6-A", "subject": "Science",
        "duration_minutes": 45, "instructions": "Answer all questions.", "include_answer_key": True,
        "sections": [{"title": "Section A: Questions", "questions": [{"id": q["id"], "marks": q["marks"]} for q in questions]}],
        **over,
    }


def test_paper_docx_has_header_marks_options_and_answer_key(client, provider, note_id):
    qs = make_set(client, provider, note_id, count=5).json()
    accept_all(client, qs["id"])
    questions = qs["questions"]
    questions[0]["marks"] = 3  # the teacher changes marks on the paper
    resp = client.post("/api/papers/docx", json=_paper_body(questions), headers=teacher(client))
    assert resp.status_code == 200 and 'filename="Unit_Test_Magnets.docx"' in resp.headers["content-disposition"]

    doc = Document(io.BytesIO(resp.content))
    lines = [p.text for p in doc.paragraphs]
    total = sum(q["marks"] for q in questions)
    text = "\n".join(lines)
    assert lines[0] == "myClassBoard Demo School" and lines[1] == "Unit Test: Magnets"
    assert f"Maximum marks: {total:g}" in text and "Time: 45 minutes" in text and "Class: 6-A" in text
    assert f"Section A: Questions   [{total:g} marks]" in text
    assert "Answer Key" in text
    for i, q in enumerate(questions, 1):
        assert f"{i}. {q['text']}" in text
    mcq = next(q for q in questions if q["type"] == "mcq")
    assert "(a) " + mcq["options"][0] in text  # options are lettered
    letter = "abcd"[mcq["options"].index(mcq["answer"])]
    assert f"({letter}) {mcq['answer']}" in text  # and the key names the right letter


def test_paper_without_answer_key_and_with_several_sections(client, provider, note_id):
    qs = make_set(client, provider, note_id, count=5).json()
    accept_all(client, qs["id"])
    q = qs["questions"]
    body = _paper_body(q, include_answer_key=False)
    body["sections"] = [{"title": "Part 1", "questions": [{"id": q[0]["id"], "marks": 1}]}, {"title": "Part 2", "questions": [{"id": x["id"], "marks": 2} for x in q[1:]]}]
    doc = Document(io.BytesIO(client.post("/api/papers/docx", json=body, headers=teacher(client)).content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Part 1   [1 marks]" in text and "Part 2   [8 marks]" in text and "Answer Key" not in text


def test_paper_only_accepts_accepted_questions_and_valid_input(client, provider, note_id):
    qs = make_set(client, provider, note_id).json()
    headers = teacher(client)
    draft = client.post("/api/papers/docx", json=_paper_body(qs["questions"][:2]), headers=headers)
    assert draft.status_code == 422 and "accepted" in draft.json()["detail"]
    accept_all(client, qs["id"])
    q0 = qs["questions"][0]
    twice = _paper_body([q0, q0])
    assert client.post("/api/papers/docx", json=twice, headers=headers).status_code == 422
    assert client.post("/api/papers/docx", json=_paper_body([{"id": 999999, "marks": 1}]), headers=headers).status_code == 404
    assert client.post("/api/papers/docx", json={"title": "x", "sections": []}, headers=headers).status_code == 422
    assert client.post("/api/papers/docx", json=_paper_body([q0], duration_minutes=0), headers=headers).status_code == 422
    empty_section = {"title": "x", "sections": [{"title": "A", "questions": []}]}
    assert client.post("/api/papers/docx", json=empty_section, headers=headers).status_code == 422


# ---------------------------------------------------------------- assessments


def test_assign_a_test_to_a_class(client, provider, note_id, db):
    qs = make_set(client, provider, note_id, count=5).json()
    accept_all(client, qs["id"])
    class_id = db.scalar(select(SchoolClass.id).where(SchoolClass.name == "6-A"))
    items = [{"question_id": q["id"], "marks": 2} for q in qs["questions"]]
    resp = client.post("/api/assessments", json={"title": "Magnets quiz", "class_id": class_id, "items": items}, headers=teacher(client))
    assert resp.status_code == 201
    assert (resp.json()["class_name"], resp.json()["question_count"], resp.json()["total_marks"]) == ("6-A", 5, 10)
    assert [a["title"] for a in client.get("/api/assessments", headers=teacher(client)).json()] == ["Magnets quiz"]


def test_assessment_rules(client, provider, note_id, db):
    qs = make_set(client, provider, note_id, count=5).json()
    headers = teacher(client)
    class_id = db.scalar(select(SchoolClass.id).where(SchoolClass.name == "6-A"))
    items = [{"question_id": qs["questions"][0]["id"], "marks": 1}]
    assert client.post("/api/assessments", json={"title": "T", "class_id": class_id, "items": items}, headers=headers).status_code == 422  # not accepted yet
    accept_all(client, qs["id"])
    assert client.post("/api/assessments", json={"title": "T", "class_id": 9999, "items": items}, headers=headers).status_code == 422
    assert client.post("/api/assessments", json={"title": "T", "class_id": class_id, "items": []}, headers=headers).status_code == 422
    assert client.post("/api/assessments", json={"title": "T", "class_id": class_id, "items": items * 2}, headers=headers).status_code == 422
    assert client.get("/api/assessments", headers=login(client, AARAV)).status_code == 403
