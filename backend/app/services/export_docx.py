"""DOCX export."""

import io
import re

from docx import Document


def slug(text: str, fallback: str = "note") -> str:
    """A safe filename stem."""
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")[:60] or fallback


def note_to_docx(title: str, text: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line.strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
