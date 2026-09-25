"""Printable question paper (DOCX) with an optional answer key."""

import io
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models import Question

LETTERS = "abcd"


@dataclass
class PaperItem:
    question: Question
    marks: float


@dataclass
class PaperSection:
    title: str
    items: list[PaperItem]


def fmt_marks(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:g}"


def _centered(doc: Document, text: str, size: int, bold: bool = True) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold, run.font.size = bold, Pt(size)


def build_paper_docx(
    *,
    school: str,
    title: str,
    class_name: str,
    subject: str,
    duration_minutes: int | None,
    instructions: str,
    sections: list[PaperSection],
    include_answer_key: bool,
) -> bytes:
    total = sum(item.marks for s in sections for item in s.items)
    doc = Document()

    if school.strip():
        _centered(doc, school.strip(), 16)
    _centered(doc, title.strip(), 14)
    details = [f"Class: {class_name}" if class_name else "", f"Subject: {subject}" if subject else ""]
    if duration_minutes:
        details.append(f"Time: {duration_minutes} minutes")
    details.append(f"Maximum marks: {fmt_marks(total)}")
    _centered(doc, "     |     ".join(d for d in details if d), 10, bold=False)

    if instructions.strip():
        p = doc.add_paragraph()
        p.add_run("Instructions: ").bold = True
        p.add_run(instructions.strip())

    number = 0
    for section in sections:
        if not section.items:
            continue
        section_total = sum(i.marks for i in section.items)
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        run = heading.add_run(f"{section.title}   [{fmt_marks(section_total)} marks]")
        run.bold, run.font.size = True, Pt(12)
        for item in section.items:
            number += 1
            q = item.question
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.add_run(f"{number}. {q.text}")
            p.add_run(f"   [{fmt_marks(item.marks)}]").italic = True
            if q.type == "mcq":
                for letter, option in zip(LETTERS, q.options or []):
                    doc.add_paragraph(f"({letter}) {option}").paragraph_format.left_indent = Pt(28)
            elif q.type == "true_false":
                doc.add_paragraph("True / False").paragraph_format.left_indent = Pt(28)
            elif q.type == "short":
                for _ in range(2):
                    doc.add_paragraph("_" * 78).paragraph_format.left_indent = Pt(14)
            elif q.type == "long":
                for _ in range(6):
                    doc.add_paragraph("_" * 78).paragraph_format.left_indent = Pt(14)

    if include_answer_key:
        doc.add_page_break()
        _centered(doc, "Answer Key", 14)
        number = 0
        for section in sections:
            if not section.items:
                continue
            doc.add_paragraph().add_run(section.title).bold = True
            for item in section.items:
                number += 1
                q = item.question
                answer = q.answer
                if q.type == "mcq" and q.options and q.answer in q.options:
                    answer = f"({LETTERS[q.options.index(q.answer)]}) {q.answer}"
                p = doc.add_paragraph()
                p.add_run(f"{number}. {answer}").bold = True
                if q.explanation:
                    p.add_run(f"\n{q.explanation}").italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
