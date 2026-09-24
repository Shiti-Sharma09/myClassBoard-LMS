"""Pull text out of typed documents (PDF, DOCX, TXT)."""

import io
from pathlib import Path

import pymupdf
from docx import Document

MAX_CHARS = 200_000
ALLOWED = {".pdf", ".docx", ".txt"}


class ExtractionError(Exception):
    """Raised with a message that is safe to show to the user."""


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ExtractionError("Please upload a PDF, DOCX or TXT file.")
    try:
        text = {".pdf": _pdf, ".docx": _docx, ".txt": _txt}[suffix](data)
    except ExtractionError:
        raise
    except Exception:
        raise ExtractionError("That file couldn't be read. It may be damaged or password protected.") from None
    text = text.strip()
    if not text:
        raise ExtractionError("No text was found in that file.")
    return text[:MAX_CHARS]


def _pdf(data: bytes) -> str:
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        text = "\n\n".join(page.get_text().strip() for page in doc)
        pages = doc.page_count
    if len(text.strip()) < 30 * pages:
        raise ExtractionError(
            "This PDF has no selectable text, so it looks scanned or handwritten. "
            "Use “Digitise handwriting” instead."
        )
    return text


def _docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")
